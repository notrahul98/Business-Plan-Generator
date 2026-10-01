import pytest
from openpyxl import load_workbook

from kns_plan.excel import PlanInvalid, generate, read_inputs
from kns_plan.schema import CostLine, LineType


def test_generated_formulas_match_hand_worked_figures(reference_plan, tmp_path):
    result = generate(reference_plan, tmp_path / "plan.xlsx")
    assert result.ok, result.issues
    values = {e.label: e.value for e in result.expected}
    assert values["coefficient of 'Default'"] == pytest.approx(4.098462, abs=1e-6)
    assert values["retail price of 'Eye serum stick'"] == 597_000
    assert values["brand price to ask for of 'Eye serum stick'"] == pytest.approx(8.714071, abs=1e-6)


def test_any_stack_shape_generates_and_cross_checks(mixed_plan, tmp_path):
    result = generate(mixed_plan, tmp_path / "plan.xlsx")
    assert result.ok, result.issues
    wb = load_workbook(result.path)
    assert wb.sheetnames[:3] == ["Assumptions", "Price Structure", "Retail Prices"]
    assert "Channels" not in wb.sheetnames  # modules appear only when the plan uses them
    assert {"_Inputs Settings", "_Inputs Cost Stack", "_Inputs Products"} <= set(wb.sheetnames)
    # every price is a formula, not a pasted number
    assert str(wb["Retail Prices"]["H5"].value).startswith("=")


def test_inputs_round_trip(mixed_plan, tmp_path):
    result = generate(mixed_plan, tmp_path / "plan.xlsx", audit=False)
    assert read_inputs(result.path) == mixed_plan


def test_edits_in_inputs_sheet_survive_reimport(reference_plan, tmp_path):
    path = generate(reference_plan, tmp_path / "plan.xlsx", audit=False).path
    wb = load_workbook(path)
    ws = wb["_Inputs Cost Stack"]
    row = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 4).value == "retailer")
    ws.cell(row, 7).value = 0.40
    wb.save(path)
    plan = read_inputs(path)
    assert next(line for line in plan.stacks[0].lines if line.id == "retailer").rate == 0.40


def test_invalid_plan_is_refused(reference_plan, tmp_path):
    reference_plan.stacks[0].lines.insert(
        5, CostLine(id="extra", label="Extra margin", type=LineType.PCT_OF_RESULT, rate=0.6))
    with pytest.raises(PlanInvalid):
        generate(reference_plan, tmp_path / "plan.xlsx")


def test_full_plan_every_module_cross_checks(full_plan, tmp_path):
    result = generate(full_plan, tmp_path / "plan.xlsx")
    assert result.ok, [str(i) for i in result.issues if i.level == "error"]
    for sheet in ("Assumptions", "Summary", "Competitor Analysis", "Channels", "Volume Build", "Forecast Y1",
                  "Forecast Y2", "Year Analysis", "Stock & Purchase", "Treatment Costing", "Container Freight",
                  "Market Size"):
        assert sheet in result.sheets
    assert result.sheets[:2] == ["Assumptions", "Summary"]
    assert len(result.expected) > 60


def test_full_plan_round_trip(full_plan, tmp_path):
    result = generate(full_plan, tmp_path / "plan.xlsx", audit=False)
    assert read_inputs(result.path) == full_plan
