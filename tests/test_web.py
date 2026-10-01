"""End-to-end: a user builds a plan for a new brand in the app, from login to downloading the workbook."""

import io
import re

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from kns_plan.web.app import create_app
from kns_plan.web.auth import hash_password

PASSWORD = "office-password-1"


@pytest.fixture
def client(tmp_path):
    app = create_app(tmp_path, hash_password(PASSWORD), https=False)
    with TestClient(app) as c:
        yield c


def csrf(html: str) -> str:
    return re.search(r'name="csrf" value="([^"]+)"', html).group(1)


def login(c) -> str:
    r = c.post("/login", data={"password": PASSWORD, "name": "Tester"})
    assert r.status_code == 200 and "Business plans" in r.text
    return csrf(r.text)


def grid_form(token, grid, rows, extra=None):
    data = {"csrf": token}
    for i, row in enumerate(rows):
        for k, v in row.items():
            data[f"{grid}-{i}-{k}"] = v
    data.update(extra or {})
    return data


def test_lockout_after_five_wrong_passwords(client):
    assert client.get("/", follow_redirects=False).headers["location"] == "/login"
    for _ in range(5):
        assert client.post("/login", data={"password": "nope"}).status_code == 401
    r = client.post("/login", data={"password": PASSWORD})
    assert r.status_code == 401 and "Too many wrong passwords" in r.text


def test_forms_need_the_session_token(client):
    login(client)
    r = client.post("/plans", data={"brand": "X", "fx_rate": "1", "csrf": "forged"})
    assert r.status_code == 400


def test_full_journey_for_a_new_brand(client):
    token = login(client)
    r = client.post("/plans", data={"csrf": token, "brand": "Glowlab", "currency": "EUR", "fx_rate": "18.500",
                                    "name": "Glowlab 2027"})
    assert r.status_code == 200 and "Brand &amp; market" in r.text
    pid = int(re.search(r"/plans/(\d+)/setup", str(r.url)).group(1))

    # 1. settings: two-year plan from April, IDR-style thousands in the exchange rate
    r = client.post(f"/plans/{pid}/setup", data={
        "csrf": token, "brand": "Glowlab", "origin": "France", "incoterm": "EXW Paris", "hs_code": "3304.99",
        "brand_currency": "eur", "fx_rate": "18.500", "vat_rate": "11", "rounding_step": "500",
        "start_month": "4", "start_year": "2027", "years": "2", "stock_cover_months": "2", "next": "pricing"})
    assert "Cost stack" in r.text and "Saved (version" in r.text

    # 2. cost stack page round-trips unchanged (no new version) and shows the coefficient
    page = client.get(f"/plans/{pid}/pricing").text
    assert "4.0985" in page  # example stack: no company default in this data folder

    # 3. upload a messy product file, map it, import
    wb = Workbook()
    wb.active.append(["Code", "Product name", "Content", "EXW (EUR)", "Texture"])
    wb.active.append(["GL-1", "Glow serum", "30ml", "€12,40", "gel"])
    wb.active.append(["GL-2", "Glow cream", "50 ml", "15.80", "cream"])
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post(f"/plans/{pid}/upload/products", data={"csrf": token},
                    files={"file": ("glowlab.xlsx", buf.getvalue())})
    assert "match the columns" in r.text
    upload_token = re.search(r"/import/products/([0-9a-f]+)", r.text).group(1)
    r = client.post(f"/plans/{pid}/import/products/{upload_token}", data={
        "csrf": token, "map-id": "Code", "map-name": "Product name", "map-size": "Content", "map-fob": "EXW (EUR)",
        "mode": "append"})
    assert "Imported 2 product(s)" in r.text
    assert "Glow serum" in r.text

    # 4. competitors, then a pricing rule on the serum
    r = client.post(f"/plans/{pid}/competitors", data=grid_form(token, "competitors", [
        {"sku": "gl_1", "brand": "Brand X", "product": "Night serum", "size": "50", "unit": "ml",
         "price_idr": "1.650.000"}]))
    assert r.status_code == 200 and "Saved" in r.text
    products = client.get(f"/plans/{pid}/products").text
    comp_id = re.search(r'<option value="(brand_x[^"]*)"', products).group(1)
    r = client.post(f"/plans/{pid}/products", data=grid_form(token, "skus", [
        {"id": "gl_1", "name": "Glow serum", "size": "30", "unit": "ml", "fob": "12.4", "rule_competitor": comp_id,
         "rule_discount": "20", "stack": "default"},
        {"id": "gl_2", "name": "Glow cream", "size": "50", "unit": "ml", "fob": "15.8", "stack": "default"}]))
    assert "Saved" in r.text
    assert "792,000" in r.text  # 1,650,000 / 50 × 30 = 990,000, less 20% = 792,000

    # 5. channels with a fee-based marketplace, list everything
    r = client.post(f"/plans/{pid}/channels", data=grid_form(token, "channels", [
        {"name": "Sociolla", "kind": "retail", "margin": "36"},
        {"name": "Shopee", "kind": "marketplace"}]))
    assert "Saved" in r.text
    data = grid_form(token, "listings", [{"id": "gl_1", "sociolla": "1", "shopee": "1"},
                                         {"id": "gl_2", "sociolla": "1"}])
    data.update(grid_form(token, "channels", [{"id": "sociolla", "name": "Sociolla", "kind": "retail", "margin": "36"},
                                              {"id": "shopee", "name": "Shopee", "kind": "marketplace"}]))
    data.update(grid_form(token, "fees", [{"channel": "shopee", "label": "Revenue share", "rate": "12",
                                           "plus_vat": "1"}]))
    r = client.post(f"/plans/{pid}/channels", data=data)
    assert "13.3%" in r.text  # 12% × 1.11

    # 6. drivers: preset factors, then fill them in
    for ch, preset in (("sociolla", "retail_sellthrough"), ("shopee", "marketplace")):
        client.post(f"/plans/{pid}/volume/add-driver", data={"csrf": token, "channel": ch, "preset": preset})
    page = client.get(f"/plans/{pid}/volume").text
    ids = re.findall(r'name="driver-([a-z0-9_]+)-name"', page)
    assert len(ids) == 2
    data = {"csrf": token, f"driver-{ids[0]}-name": "Stores"}
    data.update(grid_form(token, f"factors_{ids[0]}", [
        {"name": "Stores", "kind": "monthly", "m1": "2", "m2": "4", "m3": "6"},
        {"name": "Units per product per store", "kind": "constant", "value": "5"},
        {"name": "Probability", "kind": "constant", "value": "0.8"}]))
    data.update(grid_form(token, f"factors_{ids[1]}", [
        {"name": "Orders per month", "kind": "monthly", "m1": "100", "m2": "150"},
        {"name": "Units per order", "kind": "constant", "value": "1"}]))
    data.update(grid_form(token, "overrides", [{"channel": "shopee", "sku": "gl_2", "month": "3", "units": "40"}]))
    r = client.post(f"/plans/{pid}/volume", data=data)
    assert "Saved" in r.text and "Preview" in r.text

    # 7–8. generate, check, download
    r = client.post(f"/plans/{pid}/generate", data={"csrf": token})
    assert "Workbook ready" in r.text, r.text[:2000]
    fname = re.search(r'/download/([^"]+\.xlsx)"', r.text).group(1)
    xl = client.get(f"/plans/{pid}/download/{fname}")
    assert xl.status_code == 200
    wb = load_workbook(io.BytesIO(xl.content))
    assert {"Summary", "Competitor Analysis", "Forecast Y2", "Stock & Purchase"} <= set(wb.sheetnames)
    assert wb["_Inputs Settings"]["B6"].value == 18500  # fx_rate

    # edit the workbook's inputs in Excel and bring them back
    wb["_Inputs Settings"]["B6"].value = 19000
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post(f"/plans/{pid}/reimport", data={"csrf": token}, files={"file": ("edited.xlsx", buf.getvalue())})
    assert "Inputs re-imported" in r.text
    assert 'value="19000"' in client.get(f"/plans/{pid}/setup").text

    # history keeps every version and can restore
    hist = client.get(f"/plans/{pid}/versions").text
    assert "re-imported edited.xlsx" in hist
    r = client.post(f"/plans/{pid}/restore/2", data={"csrf": token})
    assert "Version 2 restored" in r.text


def test_invalid_input_is_not_saved_and_is_shown_back(client):
    token = login(client)
    r = client.post("/plans", data={"csrf": token, "brand": "B", "fx_rate": "15000"})
    pid = int(re.search(r"/plans/(\d+)/setup", str(r.url)).group(1))
    r = client.post(f"/plans/{pid}/products", data=grid_form(token, "skus", [
        {"name": "Thing", "fob": "-3", "stack": "default"}]))
    assert r.status_code == 422 and "Not saved" in r.text and 'value="-3"' in r.text


def test_default_stack_button_and_plan_import(client):
    token = login(client)
    r = client.post("/plans", data={"csrf": token, "brand": "B", "fx_rate": "15000"})
    pid = int(re.search(r"/plans/(\d+)/setup", str(r.url)).group(1))
    assert "example rates" in client.get(f"/plans/{pid}/pricing").text
    # change the distributor margin, save, make it the company default
    page = client.get(f"/plans/{pid}/pricing").text
    data = {"csrf": token, "stack-0-name": "Company default", "stack-0-cogs": "landed"}
    for name, value in re.findall(r'name="(lines_0-\d+-[a-z_]+)" value="([^"]*)"', page):
        data[name] = value
    for name in re.findall(r'<select name="(lines_0-\d+-[a-z_]+)"', page):
        m = re.search(rf'<select name="{name}">.*?<option value="([^"]*)" selected', page, re.S)
        data[name] = m.group(1) if m else ""
    row = next(n.split("-")[1] for n, v in data.items() if n.endswith("-label") and v == "Distributor margin")
    data[f"lines_0-{row}-rate"] = "33"
    assert "Saved" in client.post(f"/plans/{pid}/pricing", data=data).text
    r = client.post(f"/plans/{pid}/pricing/make-default/0", data={"csrf": token})
    assert "starting cost stack for new plans" in r.text
    r = client.post("/plans", data={"csrf": token, "brand": "C", "fx_rate": "15000"})
    assert 'value="33"' in client.get(str(r.url).replace("/setup", "/pricing")).text

    # a downloaded plan file comes back as a new plan
    exported = client.get(f"/plans/{pid}/export.json").content
    r = client.post("/plans/import", data={"csrf": token}, files={"file": ("plan.json", exported)})
    assert "Imported plan.json" in r.text and "B business plan" in r.text
    r = client.post("/plans/import", data={"csrf": token}, files={"file": ("bad.json", b"{\"x\": 1}")})
    assert "not a plan file" in r.text
