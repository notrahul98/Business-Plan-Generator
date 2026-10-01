"""Recalculate a generated workbook outside Excel, so it can be checked before download.

openpyxl writes formulas without results. LibreOffice (headless) recalculates them and saves the
values into the file, so it also opens with numbers in email previews and on phones. When
LibreOffice is not installed (e.g. a development laptop), the pure-Python `formulas` package is
used for the check only, and the downloaded file keeps no cached values.
"""

from __future__ import annotations

import os
import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

EXCEL_ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!")


def find_soffice() -> str | None:
    candidates = [os.environ.get("KNS_SOFFICE"), shutil.which("soffice"), shutil.which("soffice.exe"),
                  r"C:\Program Files\LibreOffice\program\soffice.exe",
                  r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
                  "/usr/bin/soffice", "/Applications/LibreOffice.app/Contents/MacOS/soffice"]
    return next((c for c in candidates if c and Path(c).exists()), None)


@dataclass
class Recalculated:
    backend: str
    values: dict[tuple[str, str], object]  # (sheet, coord) → value, formula cells only
    recalculated_file: Path | None = None   # LibreOffice output with cached values

    def get(self, sheet: str, coord: str):
        return self.values.get((sheet, coord))

    def errors(self) -> list[tuple[str, str, str]]:
        return [(s, c, str(v)) for (s, c), v in self.values.items() if str(v) in EXCEL_ERRORS]


def recalculate(path: Path) -> Recalculated:
    """LibreOffice when installed; the Python `formulas` package if it is missing or fails."""
    soffice = find_soffice()
    if soffice:
        try:
            result = _libreoffice(Path(path), soffice)
            if any(v is not None for v in result.values.values()):
                return result
            note = "LibreOffice returned no values"
        except (subprocess.SubprocessError, OSError, KeyError) as e:
            note = f"LibreOffice failed ({type(e).__name__})"
        fallback = _formulas(Path(path))
        fallback.backend = f"formulas ({note})"
        return fallback
    return _formulas(Path(path))


def _formula_cells(path: Path) -> list[tuple[str, str]]:
    wb = load_workbook(path)
    return [(ws.title, c.coordinate) for ws in wb.worksheets for row in ws.iter_rows() for c in row
            if isinstance(c.value, str) and c.value.startswith("=")]


# LibreOffice normally trusts the values cached in an .xlsx; ours has none, so make it recalculate
# everything on load (0 = always) via the throwaway profile used for each run.
_RECALC_ALWAYS = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema"
 xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="ODFRecalcMode" oor:op="fuse"><value>0</value></prop></item>
</oor:items>
"""


def _libreoffice(path: Path, soffice: str) -> Recalculated:
    out_dir = Path(tempfile.mkdtemp(prefix="kns_recalc_"))
    profile = Path(tempfile.mkdtemp(prefix="kns_lo_profile_"))
    (profile / "user").mkdir()
    (profile / "user" / "registrymodifications.xcu").write_text(_RECALC_ALWAYS, encoding="utf-8")
    try:
        subprocess.run([soffice, f"-env:UserInstallation={profile.as_uri()}", "--headless", "--norestore",
                        "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(out_dir),
                        str(path)], check=True, capture_output=True, timeout=180)
        result = out_dir / path.name
        wb = load_workbook(result, data_only=True)
        values = {(s, c): wb[s][c].value for s, c in _formula_cells(path)}
    finally:
        shutil.rmtree(profile, ignore_errors=True)
    return Recalculated("libreoffice", values, result)


def _formulas(path: Path) -> Recalculated:
    os.environ.setdefault("TQDM_DISABLE", "1")  # formulas prints progress bars otherwise
    import formulas  # dev fallback

    solution = formulas.ExcelModel().loads(str(path)).finish().calculate()
    by_key = {}
    for key, ranges in solution.items():
        # keys look like "'[book.xlsx]SHEET NAME'!E12"
        if "!" not in key:
            continue
        sheet_part, coord = key.rsplit("!", 1)
        sheet = sheet_part.strip("'").split("]", 1)[-1].upper()
        value = ranges.value[0][0] if hasattr(ranges, "value") else ranges
        by_key[(sheet, coord.replace("$", "").upper())] = value
    values = {}
    for s, c in _formula_cells(path):
        v = by_key.get((s.upper(), c))
        if v is not None and not isinstance(v, (int, float, str, bool)):
            v = str(v)  # XlError and similar
        values[(s, c)] = v
    return Recalculated("formulas", values)
