"""Assemble the workbook from the modules the plan uses, then check it before it is handed out."""

from __future__ import annotations

import math
import shutil
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook

from ..recalc import recalculate
from ..schema import Plan
from ..validate import Issue, check_plan, errors
from . import addons, forecast, price
from .inputs import write_inputs
from .price import Expected


class PlanInvalid(Exception):
    def __init__(self, issues: list[Issue]):
        self.issues = issues
        super().__init__("\n".join(str(i) for i in issues))


@dataclass
class GenerateResult:
    path: Path
    issues: list[Issue] = field(default_factory=list)  # validation warnings plus audit findings
    expected: list[Expected] = field(default_factory=list)
    audit_backend: str | None = None
    sheets: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not errors(self.issues)


def build(plan: Plan) -> tuple[Workbook, list[Expected]]:
    wb = Workbook()
    wb.remove(wb.active)
    expected: list[Expected] = []
    refs = write_inputs(wb, plan)
    inputs = list(wb.sheetnames)

    stacks = price.write_price_structure(wb, plan, refs, expected)
    matched = price.write_competitors(wb, plan, refs, expected)
    skus = price.write_retail_prices(wb, plan, refs, stacks, matched, expected)
    if plan.channels:
        margins, weights = forecast.write_channels(wb, plan, refs, skus, expected)
        if plan.drivers:
            totals = forecast.write_volume_build(wb, plan, refs)
            rows = forecast.write_forecasts(wb, plan, refs, skus, weights, totals, expected)
            year_totals = forecast.write_year_analysis(wb, plan, refs, margins, rows, expected)
            forecast.write_summary(wb, plan, refs, year_totals, expected)
            forecast.write_stock(wb, plan, refs, skus, rows, expected)
    if plan.packages:
        addons.write_treatments(wb, plan, refs, skus, expected)
    if plan.containers:
        addons.write_freight(wb, plan, refs, skus, expected)
    if plan.market:
        addons.write_market(wb, plan)

    outputs = [name for name in wb.sheetnames if name not in inputs]
    addons.write_assumptions(wb, plan, refs, [addons.ASSUMPTIONS_SHEET] + outputs)
    front = [addons.ASSUMPTIONS_SHEET] + ([forecast.SUMMARY_SHEET] if forecast.SUMMARY_SHEET in outputs else [])
    order = front + [n for n in outputs if n not in front] + inputs
    wb._sheets = [wb[name] for name in order]
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True  # the file carries formulas only; Excel computes them on opening
    return wb, expected


def generate(plan: Plan, path: Path | str, audit: bool = True) -> GenerateResult:
    issues = check_plan(plan)
    if errors(issues):
        raise PlanInvalid(issues)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb, expected = build(plan)
    wb.save(path)
    result = GenerateResult(path=path, issues=list(issues), expected=expected, sheets=list(wb.sheetnames))
    if audit:
        _audit(result)
    return result


def _audit(result: GenerateResult) -> None:
    recalc = recalculate(result.path)
    result.audit_backend = recalc.backend
    for sheet, coord, err in recalc.errors():
        result.issues.append(Issue("error", f"{sheet}!{coord}", f"formula returns {err}"))
    for e in result.expected:
        got = recalc.get(e.sheet, e.coord)
        if not isinstance(got, (int, float)) or not math.isclose(got, e.value, rel_tol=1e-9, abs_tol=1e-6):
            result.issues.append(Issue("error", f"{e.sheet}!{e.coord}",
                                       f"{e.label}: workbook gives {got!r}, engine gives {e.value!r}"))
    # The download stays the file this app wrote (LibreOffice's re-saved copy is only used for the
    # check); Excel recalculates it on opening.
    if recalc.recalculated_file:
        shutil.rmtree(recalc.recalculated_file.parent, ignore_errors=True)
    if recalc.backend == "formulas":
        result.issues.append(Issue("warning", "audit", "LibreOffice is not installed here; checked with the "
                                                       "Python fallback instead"))
    elif recalc.backend != "libreoffice":
        result.issues.append(Issue("warning", "audit", f"checked with the Python fallback: {recalc.backend}"))
