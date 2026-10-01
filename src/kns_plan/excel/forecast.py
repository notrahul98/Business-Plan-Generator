"""Channels, Volume Build, Forecast Y1…Yn, Year Analysis, Summary and Stock & Purchase sheets."""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from .. import model
from ..schema import Plan
from .inputs import InputRefs, month_labels
from .price import Expected, SkuRefs
from .style import BOLD, IDR, INPUT, MONEY, PCT, SUBTITLE, SUBTOTAL_FILL, TITLE, TOP_RULE, header_row, ref, widths

CHANNELS_SHEET = "Channels"
VOLUME_SHEET = "Volume Build"
ANALYSIS_SHEET = "Year Analysis"
SUMMARY_SHEET = "Summary"
STOCK_SHEET = "Stock & Purchase"
UNITS = "#,##0"
MONTH_COL0 = 4  # Volume Build: first month in column D


def forecast_sheet(year: int) -> str:
    return f"Forecast Y{year}"


# ------------------------------------------------------------------ channels

def write_channels(wb: Workbook, plan: Plan, refs: InputRefs, skus: dict[str, SkuRefs],
                   expected: list[Expected]) -> tuple[dict[str, str], dict[tuple[str, str], str]]:
    """Returns channel → margin cell, and (channel, sku) → weight cell."""
    ws = wb.create_sheet(CHANNELS_SHEET)
    ws.cell(1, 2, f"Channels and product listing – {plan.settings.brand}").font = TITLE
    ws.cell(2, 2, "Weight = units of the product per unit of the channel's volume drivers (blank = not listed). "
                  "Treatment channels take weights from their package: usage ÷ product size.").font = SUBTITLE
    header_row(ws, 4, ["Product", "Forecast price (IDR)"] + [ch.name for ch in plan.channels], start_col=2)
    for j, ch in enumerate(plan.channels, start=4):
        ws.cell(4, j, f"={refs.channels[ch.id]['name']}")
    ws.cell(5, 2, "Channel margin").font = BOLD
    margins: dict[str, str] = {}
    for j, ch in enumerate(plan.channels, start=4):
        col = get_column_letter(j)
        fees = refs.fees.get(ch.id, [])
        if fees:
            parts = [f'{f["rate"]}*(1+IF(LOWER(TRIM({f["plus_vat"]}))="yes",{refs.vat},0))' for f in fees]
            formula = "=" + "+".join(parts)
        else:
            m = refs.channels[ch.id]["margin"]
            formula = f'=IF({m}="",0,{m})'
        c = ws.cell(5, j, formula)
        c.number_format, c.font, c.fill = PCT, BOLD, SUBTOTAL_FILL
        margins[ch.id] = ref(CHANNELS_SHEET, f"{col}5")
        expected.append(Expected(CHANNELS_SHEET, f"{col}5", model.effective_margin(plan, ch),
                                 f"margin of channel '{ch.name}'"))

    weights: dict[tuple[str, str], str] = {}
    for i, sku in enumerate(plan.skus):
        r = 6 + i
        ws.cell(r, 2, f"={skus[sku.id].name}")
        ws.cell(r, 3, f"={skus[sku.id].forecast_price}").number_format = IDR
        for j, ch in enumerate(plan.channels, start=4):
            col = get_column_letter(j)
            if ch.package:
                usage = [u for (p, s), u in refs.package_usage.items() if p == ch.package and s == sku.id]
                if not usage:
                    continue
                size = refs.skus[sku.id]["size"]
                formula = f'=({"+".join(usage)})/IF({size}="",1,{size})'
            else:
                lst = refs.listings[(ch.id, sku.id)]
                formula = f'=IF({lst}="","",{lst})'
            c = ws.cell(r, j, formula)
            c.number_format = "0.####"
            if (ch.id, sku.id) in _weighted(plan):
                weights[(ch.id, sku.id)] = ref(CHANNELS_SHEET, f"{col}{r}")
    widths(ws, {"A": 2, "B": 42, "C": 16, **{get_column_letter(j): 16 for j in range(4, 4 + len(plan.channels))}})
    ws.freeze_panes = "D6"
    return margins, weights


def _weighted(plan: Plan) -> set[tuple[str, str]]:
    return {(ch.id, s) for ch in plan.channels for s in model.channel_weights(plan, ch)}


# ------------------------------------------------------------------ volume build

def write_volume_build(wb: Workbook, plan: Plan, refs: InputRefs) -> dict[str, int]:
    """One block per driver: its factors by month and their product. Returns channel → total row."""
    ws = wb.create_sheet(VOLUME_SHEET)
    n = plan.settings.months
    labels = month_labels(plan.settings)
    ws.cell(1, 2, "Volume build").font = TITLE
    ws.cell(2, 2, "Each driver's monthly volume = its factors multiplied together. A channel's volume is the sum of "
                  "its drivers; product units = channel volume × product weight (see Forecast sheets).").font = SUBTITLE
    header_row(ws, 4, ["Driver / factor", "Unit"] + labels, start_col=2)
    row = 5
    totals: dict[str, int] = {}
    names = {ch.id: ch.name for ch in plan.channels}
    for ch in plan.channels:
        driver_rows = []
        for d in (d for d in plan.drivers if d.channel == ch.id):
            ws.cell(row, 2, f"{names[ch.id]} – {d.name}").font = BOLD
            row += 1
            first = row
            for f in refs.factors[d.id]:
                ws.cell(row, 2, f["name"])
                ws.cell(row, 3, next(x.unit for x in d.factors if x.name == f["name"]))
                for m in range(n):
                    c = ws.cell(row, MONTH_COL0 + m, f'=IF({f["kind"]}="constant",{f["constant"]},{f["months"][m]})')
                    c.number_format, c.fill = "#,##0.####", INPUT
                row += 1
            ws.cell(row, 2, "Driver volume")
            for m in range(n):
                col = get_column_letter(MONTH_COL0 + m)
                ws.cell(row, MONTH_COL0 + m, f"=PRODUCT({col}{first}:{col}{row - 1})").number_format = "#,##0.##"
            driver_rows.append(row)
            row += 2
        ws.cell(row, 2, f"{names[ch.id]} – total volume").font = BOLD
        for m in range(n):
            col = get_column_letter(MONTH_COL0 + m)
            c = ws.cell(row, MONTH_COL0 + m, "=" + ("+".join(f"{col}{r}" for r in driver_rows) or "0"))
            c.number_format, c.font, c.fill, c.border = "#,##0.##", BOLD, SUBTOTAL_FILL, TOP_RULE
        ws.cell(row, 2).fill = SUBTOTAL_FILL
        totals[ch.id] = row
        row += 3
    widths(ws, {"A": 2, "B": 44, "C": 14, **{get_column_letter(MONTH_COL0 + m): 10 for m in range(n)}})
    ws.freeze_panes = ws.cell(5, MONTH_COL0)
    return totals


# ------------------------------------------------------------------ forecast per year

def write_forecasts(wb: Workbook, plan: Plan, refs: InputRefs, skus: dict[str, SkuRefs],
                    weights: dict[tuple[str, str], str], totals: dict[str, int],
                    expected: list[Expected]) -> dict[int, tuple[int, int]]:
    """Returns year → (first, last) data row on that year's sheet."""
    labels = month_labels(plan.settings)
    units = model.units(plan)
    names = {ch.id: ch.name for ch in plan.channels}
    o = refs.overrides
    pairs = [(ch.id, s.id) for ch in plan.channels for s in plan.skus if (ch.id, s.id) in weights]
    out = {}
    for y in range(1, plan.settings.years + 1):
        ws = wb.create_sheet(forecast_sheet(y))
        ws.cell(1, 2, f"Forecast year {y} – units by channel and product").font = TITLE
        ws.cell(2, 2, "Units = channel volume × product weight, unless the '_Inputs Overrides' sheet has a value "
                      "for that channel, product and month.").font = SUBTITLE
        heads = ["Channel ID", "Channel", "Product", "SKU ID", "Weight"] + labels[(y - 1) * 12: y * 12] + [
            "Total units", "Forecast price (IDR)", "Sales incl. VAT (IDR)", "Cost of goods (IDR)"]
        header_row(ws, 4, heads)
        first = 5
        for i, (ch_id, sku_id) in enumerate(pairs):
            r = first + i
            ws.cell(r, 1, ch_id)
            ws.cell(r, 2, names[ch_id])
            ws.cell(r, 3, f"={skus[sku_id].name}")
            ws.cell(r, 4, sku_id)
            w = weights[(ch_id, sku_id)]
            ws.cell(r, 5, f'=IF({w}="",0,{w})').number_format = "0.####"
            for k in range(12):
                m = (y - 1) * 12 + k + 1  # plan month index, 1-based
                vb = ref(VOLUME_SHEET, f"{get_column_letter(MONTH_COL0 + m - 1)}{totals[ch_id]}")
                crit = f'{o["channel"]},$A{r},{o["sku"]},$D{r},{o["month"]},{m}'
                ws.cell(r, 6 + k, f"=IF(COUNTIFS({crit})>0,SUMIFS({o['units']},{crit}),$E{r}*{vb})"
                        ).number_format = UNITS
            ws.cell(r, 18, f"=SUM(F{r}:Q{r})").number_format = UNITS
            ws.cell(r, 19, f"={skus[sku_id].forecast_price}").number_format = IDR
            ws.cell(r, 20, f"=R{r}*S{r}").number_format = IDR
            ws.cell(r, 21, f"=R{r}*{skus[sku_id].landed}").number_format = IDR
            expected.append(Expected(forecast_sheet(y), f"R{r}", sum(model.year_slice(units[(ch_id, sku_id)], y)),
                                     f"Y{y} units of {sku_id} in {ch_id}"))
        last = first + len(pairs) - 1
        tr = last + 1
        ws.cell(tr, 2, "Total").font = BOLD
        for col in [get_column_letter(c) for c in range(6, 19)] + ["T", "U"]:
            c = ws[f"{col}{tr}"]
            c.value, c.font, c.border = f"=SUM({col}{first}:{col}{last})", BOLD, TOP_RULE
            c.number_format = IDR if col in "TU" else UNITS
        widths(ws, {"A": 12, "B": 22, "C": 38, "D": 14, "E": 8, **{get_column_letter(c): 9 for c in range(6, 18)},
                    "R": 11, "S": 14, "T": 16, "U": 16})
        ws.column_dimensions["A"].hidden = True
        ws.column_dimensions["D"].hidden = True
        ws.freeze_panes = "F5"
        out[y] = (first, last)
    return out


# ------------------------------------------------------------------ year analysis and summary

MEASURES = [("units", "Units", UNITS), ("sales", "Retail sales incl. VAT", IDR),
            ("margin_rate", "Channel margin %", PCT), ("margin", "Channel margin", IDR),
            ("wholesale", "Wholesale incl. VAT", IDR), ("vat", "VAT", IDR),
            ("net_wholesale", "Net wholesale (KNS revenue excl. VAT)", IDR), ("cogs", "Cost of goods", IDR),
            ("gross_profit", "Gross profit", IDR)]


def write_year_analysis(wb: Workbook, plan: Plan, refs: InputRefs, margins: dict[str, str],
                        rows: dict[int, tuple[int, int]], expected: list[Expected]) -> dict[int, int]:
    """Returns year → total row on the Year Analysis sheet."""
    ws = wb.create_sheet(ANALYSIS_SHEET)
    analysis = model.analyse(plan)
    ws.cell(1, 2, f"Year analysis – {plan.settings.brand}").font = TITLE
    ws.cell(2, 2, "Wholesale = retail sales − channel margin. VAT is the VAT inside wholesale "
                  "(wholesale × VAT ÷ (1 + VAT)). Gross profit = net wholesale − cost of goods.").font = SUBTITLE
    heads = ["Channel"] + [m[1] for m in MEASURES] + ["Contribution %"]
    totals: dict[int, int] = {}
    row = 4
    for y in range(1, plan.settings.years + 1):
        ws.cell(row, 2, f"Year {y}").font = BOLD
        row += 1
        header_row(ws, row, heads, start_col=2)
        row += 1
        first, last = rows[y]
        sheet = forecast_sheet(y)
        rng = {k: f"'{sheet}'!${k}${first}:${k}${last}" for k in ("A", "R", "T", "U")}
        start = row
        n = len(plan.channels)
        for ch in plan.channels:
            r = row
            ws.cell(r, 2, f"={refs.channels[ch.id]['name']}")
            crit = f'{rng["A"]},"{ch.id}"'
            ws.cell(r, 3, f"=SUMIF({crit},{rng['R']})").number_format = UNITS
            ws.cell(r, 4, f"=SUMIF({crit},{rng['T']})").number_format = IDR
            ws.cell(r, 5, f"={margins[ch.id]}").number_format = PCT
            ws.cell(r, 6, f"=D{r}*E{r}").number_format = IDR
            ws.cell(r, 7, f"=D{r}-F{r}").number_format = IDR
            ws.cell(r, 8, f"=G{r}*{refs.vat}/(1+{refs.vat})").number_format = IDR
            ws.cell(r, 9, f"=G{r}-H{r}").number_format = IDR
            ws.cell(r, 10, f"=SUMIF({crit},{rng['U']})").number_format = IDR
            ws.cell(r, 11, f"=I{r}-J{r}").number_format = IDR
            ws.cell(r, 12, f"=IF($D${start + n}=0,0,D{r}/$D${start + n})").number_format = PCT
            cy = analysis.by_year[y][ch.id]
            for col, key in (("D", "sales"), ("I", "net_wholesale"), ("K", "gross_profit")):
                expected.append(Expected(ANALYSIS_SHEET, f"{col}{r}", getattr(cy, key), f"Y{y} {key} of {ch.id}"))
            row += 1
        tr = row
        ws.cell(tr, 2, "Total").font = BOLD
        for col in "CDFGHIJK":
            c = ws[f"{col}{tr}"]
            c.value, c.font, c.border = f"=SUM({col}{start}:{col}{tr - 1})", BOLD, TOP_RULE
            c.number_format = UNITS if col == "C" else IDR
        c = ws[f"E{tr}"]
        c.value, c.number_format, c.font, c.border = f"=IF(D{tr}=0,0,F{tr}/D{tr})", PCT, BOLD, TOP_RULE
        ws[f"L{tr}"].value, ws[f"L{tr}"].number_format = f"=SUM(L{start}:L{tr - 1})", PCT
        totals[y] = tr
        row += 3
    widths(ws, {"A": 2, "B": 28, "C": 11, "D": 17, "E": 10, "F": 16, "G": 17, "H": 15, "I": 19, "J": 16,
                "K": 16, "L": 13})
    return totals


def write_summary(wb: Workbook, plan: Plan, refs: InputRefs, totals: dict[int, int],
                  expected: list[Expected]) -> None:
    ws = wb.create_sheet(SUMMARY_SHEET)
    s = plan.settings
    analysis = model.analyse(plan)
    years = list(range(1, s.years + 1))
    ws.cell(1, 2, f"Summary – {s.brand}").font = TITLE
    ws.cell(2, 2, f"All channels. Brand-currency figures at {s.fx_rate:,.2f} IDR per {s.brand_currency}, "
                  "linked to the exchange rate input.").font = SUBTITLE
    cols = {"units": "C", "sales": "D", "margin": "F", "wholesale": "G", "vat": "H", "net_wholesale": "I",
            "cogs": "J", "gross_profit": "K"}
    rows = [m for m in MEASURES if m[0] != "margin_rate"] + [("gp_pct", "Gross profit % of net wholesale", PCT)]
    for block, (title, divide) in enumerate([("IDR", None), (s.brand_currency, refs.fx)]):
        top = 4 + block * (len(rows) + 4)
        header_row(ws, top, [title] + [f"Year {y}" for y in years] + ["Total"], start_col=2)
        for i, (key, label, fmt) in enumerate(rows):
            r = top + 1 + i
            ws.cell(r, 2, label)
            for j, y in enumerate(years):
                col = get_column_letter(3 + j)
                if key == "gp_pct":
                    net_r, gp_r = top + 1 + [k for k, *_ in rows].index("net_wholesale"), r - 1
                    formula = f"=IF({col}{net_r}=0,0,{col}{gp_r}/{col}{net_r})"
                else:
                    src = ref(ANALYSIS_SHEET, f"{cols[key]}{totals[y]}")
                    formula = f"={src}" if (divide is None or key == "units") else f"={src}/{divide}"
                c = ws.cell(r, 3 + j, formula)
                c.number_format = fmt if (divide is None or fmt != IDR) else MONEY
            tcol = get_column_letter(3 + len(years))
            if key == "gp_pct":
                net_r, gp_r = top + 1 + [k for k, *_ in rows].index("net_wholesale"), r - 1
                total = f"=IF({tcol}{net_r}=0,0,{tcol}{gp_r}/{tcol}{net_r})"
            else:
                total = f"=SUM(C{r}:{get_column_letter(2 + len(years))}{r})"
            c = ws.cell(r, 3 + len(years), total)
            c.number_format, c.font = (fmt if (divide is None or fmt != IDR) else MONEY), BOLD
            if divide is None and key in ("sales", "net_wholesale", "gross_profit"):
                expected.append(Expected(SUMMARY_SHEET, f"{tcol}{r}",
                                         sum(getattr(analysis.total(y), key) for y in years), f"total {key}"))
        ws.cell(top + 1 + [k for k, *_ in rows].index("net_wholesale"), 2).font = BOLD
    widths(ws, {"A": 2, "B": 38, **{get_column_letter(3 + j): 17 for j in range(len(years) + 1)}})


# ------------------------------------------------------------------ stock

def write_stock(wb: Workbook, plan: Plan, refs: InputRefs, skus: dict[str, SkuRefs],
                rows: dict[int, tuple[int, int]], expected: list[Expected]) -> None:
    ws = wb.create_sheet(STOCK_SHEET)
    s = plan.settings
    stock = model.stock(plan)
    cover = refs.settings["stock_cover_months"]
    ws.cell(1, 2, f"Stock and purchases – {s.brand}").font = TITLE
    ws.cell(2, 2, "Purchase = sales + closing target − opening stock (never below 0), rounded up to the MOQ. "
                  "Closing target = cover months × next year's monthly sales.").font = SUBTITLE
    heads = ["Product", "Sales (units)", "Opening", "Closing target", "Purchase", "Closing", "Cover (months)",
             f"Purchase at brand price ({s.brand_currency})", "Purchase at cost of goods (IDR)", "Lead time (months)"]
    n = len(plan.skus)
    block = n + 5
    top = {y: 4 + (y - 1) * block for y in range(1, s.years + 1)}
    first_row = {y: top[y] + 2 for y in top}
    for y in top:
        ws.cell(top[y], 2, f"Year {y}").font = BOLD
        header_row(ws, top[y] + 1, heads, start_col=2)
        fs = forecast_sheet(y)
        f_first, f_last = rows[y]
        for i, sku in enumerate(plan.skus):
            r = first_row[y] + i
            nxt = first_row[y + 1] + i if y + 1 in top else r
            sr = refs.skus[sku.id]
            ws.cell(r, 2, f"={skus[sku.id].name}")
            ws.cell(r, 3, f"=SUMIF('{fs}'!$D${f_first}:$D${f_last},\"{sku.id}\",'{fs}'!$R${f_first}:$R${f_last})"
                    ).number_format = UNITS
            prev = f"G{first_row[y - 1] + i}" if y > 1 else sr["opening"]
            ws.cell(r, 4, f"={prev}" if y > 1 else f'=IF({prev}="",0,{prev})').number_format = UNITS
            ws.cell(r, 5, f"={cover}/12*C{nxt}").number_format = UNITS
            need = f"MAX(0,C{r}+E{r}-D{r})"
            ws.cell(r, 6, f'=IF({sr["moq"]}="",{need},ROUNDUP({need}/{sr["moq"]},0)*{sr["moq"]})').number_format = UNITS
            ws.cell(r, 7, f"=D{r}+F{r}-C{r}").number_format = UNITS
            ws.cell(r, 8, f'=IF(C{nxt}=0,"",G{r}/(C{nxt}/12))').number_format = "0.0"
            ws.cell(r, 9, f'=IF({skus[sku.id].cost_fob}="",0,F{r}*{skus[sku.id].cost_fob})').number_format = MONEY
            ws.cell(r, 10, f"=F{r}*{skus[sku.id].landed}").number_format = IDR
            ws.cell(r, 11, f"={sr['lead']}").number_format = "0"
            expected.append(Expected(STOCK_SHEET, f"F{r}", stock[sku.id][y - 1].purchase, f"Y{y} purchase of {sku.id}"))
            expected.append(Expected(STOCK_SHEET, f"G{r}", stock[sku.id][y - 1].closing, f"Y{y} closing of {sku.id}"))
        tr = first_row[y] + n
        ws.cell(tr, 2, "Total").font = BOLD
        for col in "CDEFGIJ":
            c = ws[f"{col}{tr}"]
            c.value, c.font, c.border = f"=SUM({col}{first_row[y]}:{col}{tr - 1})", BOLD, TOP_RULE
            c.number_format = MONEY if col == "I" else (IDR if col == "J" else UNITS)
    widths(ws, {"A": 2, "B": 42, "C": 12, "D": 11, "E": 13, "F": 11, "G": 11, "H": 13, "I": 20, "J": 22, "K": 12})
