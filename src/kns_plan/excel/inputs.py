"""The `_Inputs` sheets: structured tables that hold every input of the plan.

Output sheets only reference these cells. Re-uploading a workbook reads these sheets and
nothing else, so edits made here in Excel survive a regenerate; edits on output sheets do not.
Values change live in Excel; adding rows (products, channels, drivers…) needs a regenerate,
except overrides, which are looked up by formula and can be added directly.
"""

from __future__ import annotations

from collections.abc import Callable
from dataclasses import dataclass, field

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from ..schema import (SKU, Channel, Competitor, ContainerOption, CostLine, CostStack, Driver, Factor,
                      LineType, MarginComponent, MarketRow, Override, Package, PackageComponent, Plan,
                      PricingRule, Settings)
from .style import IDR, MONEY, PCT, header_row, ref, widths

SETTINGS_SHEET = "_Inputs Settings"
STACK_SHEET = "_Inputs Cost Stack"
PRODUCTS_SHEET = "_Inputs Products"
COMPETITORS_SHEET = "_Inputs Competitors"
CHANNELS_SHEET = "_Inputs Channels"
MARGINS_SHEET = "_Inputs Channel Fees"
LISTINGS_SHEET = "_Inputs Listings"
DRIVERS_SHEET = "_Inputs Drivers"
OVERRIDES_SHEET = "_Inputs Overrides"
PACKAGES_SHEET = "_Inputs Packages"
CONTAINERS_SHEET = "_Inputs Containers"
MARKET_SHEET = "_Inputs Market"
OVERRIDE_ROWS = 2000  # override lookups cover this many rows, so rows can be added in Excel

SETTINGS_ROWS = [  # key, description, number format
    ("brand", "Brand name", None),
    ("origin", "Country of origin", None),
    ("incoterm", "Incoterm of the brand price, e.g. FOB Korea, EXW Spain", None),
    ("brand_currency", "Currency of the brand price", None),
    ("fx_rate", "IDR per 1 unit of brand currency", MONEY),
    ("hs_code", "HS code used for import duties", None),
    ("rounding_step", "Retail prices round up to this many IDR (0, 100, 500 or 1000)", IDR),
    ("vat_rate", "VAT on sales, used in the year analysis", PCT),
    ("start_year", "First year of the plan", "0"),
    ("start_month", "First month of the plan (1 = January)", "0"),
    ("years", "Plan length in years (1 to 5)", "0"),
    ("stock_cover_months", "Closing stock target, in months of the next year's sales", "0.0"),
]
STACK_COLS = ["Stack ID", "Stack name", "Order", "Line ID", "Label", "Type", "Rate", "Base",
              "Amount", "Amount currency", "Note", "Cost of goods line"]
PRODUCT_COLS = ["SKU ID", "Name", "Size", "Size label", "Unit", "Brand price", "Target retail (IDR)",
                "Rule competitor", "Rule discount", "Cost stack", "Opening units", "MOQ", "Lead time (months)"]
COMPETITOR_COLS = ["Competitor ID", "For product", "Brand", "Product", "Origin", "Size", "Unit", "Price (IDR)",
                   "Source", "Registration", "Notes"]
CHANNEL_COLS = ["Channel ID", "Name", "Kind", "Margin", "Package"]
MARGIN_COLS = ["Channel ID", "Fee", "Rate", "Plus VAT"]
DRIVER_COLS = ["Driver ID", "Channel ID", "Driver name", "Preset", "Factor", "Kind", "Unit", "Constant"]
OVERRIDE_COLS = ["Channel ID", "SKU ID", "Month", "Units"]
PACKAGE_COLS = ["Package ID", "Package name", "Price multiplier", "Duration (min)", "SKU ID", "Usage per treatment"]
CONTAINER_COLS = ["Container", "Cost", "Units"]
MARKET_COLS = ["Segment", "Metric", "Value", "Source"]


@dataclass
class InputRefs:
    """Absolute references to input cells, used by output-sheet formulas."""
    settings: dict[str, str] = field(default_factory=dict)
    lines: dict[tuple[str, str], dict[str, str]] = field(default_factory=dict)  # (stack, line) → refs
    skus: dict[str, dict[str, str]] = field(default_factory=dict)
    competitors: dict[str, dict[str, str]] = field(default_factory=dict)
    channels: dict[str, dict[str, str]] = field(default_factory=dict)
    fees: dict[str, list[dict[str, str]]] = field(default_factory=dict)  # channel → [{rate, plus_vat}]
    listings: dict[tuple[str, str], str] = field(default_factory=dict)  # (channel, sku) → weight cell
    factors: dict[str, list[dict]] = field(default_factory=dict)  # driver → [{name, kind, constant, months}]
    overrides: dict[str, str] = field(default_factory=dict)  # column → bounded range
    package_usage: dict[tuple[str, str], str] = field(default_factory=dict)  # (package, sku) → usage cell
    packages: dict[str, dict[str, str]] = field(default_factory=dict)
    containers: list[dict[str, str]] = field(default_factory=list)

    @property
    def fx(self) -> str:
        return self.settings["fx_rate"]

    @property
    def step(self) -> str:
        return self.settings["rounding_step"]

    @property
    def vat(self) -> str:
        return self.settings["vat_rate"]


# ------------------------------------------------------------------ writing

def _table(ws, name: str, ncols: int, nrows: int) -> None:
    last = get_column_letter(ncols) + str(max(nrows, 1) + 1)
    t = Table(displayName=name, ref=f"A1:{last}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(t)


Col = tuple[str, Callable, str | None]  # header, value getter, number format


def _write_rows(wb: Workbook, sheet: str, table: str, cols: list[Col], items: list,
                col_widths: dict[str, float] | None = None) -> list[int]:
    """Write one row per item; returns the row numbers used."""
    ws = wb.create_sheet(sheet)
    header_row(ws, 1, [c[0] for c in cols])
    rows = []
    for r, item in enumerate(items, start=2):
        for j, (_, get, fmt) in enumerate(cols, start=1):
            cell = ws.cell(r, j, get(item))
            if fmt:
                cell.number_format = fmt
        rows.append(r)
    _table(ws, table, len(cols), len(items))
    if col_widths:
        widths(ws, col_widths)
    ws.freeze_panes = "A2"
    return rows


def _col(cols: list[Col], header: str) -> str:
    return get_column_letter([c[0] for c in cols].index(header) + 1)


def month_labels(settings: Settings) -> list[str]:
    names = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    out = []
    for m in range(settings.months):
        k = settings.start_month - 1 + m
        out.append(f"{names[k % 12]} {settings.start_year + k // 12}")
    return out


def write_inputs(wb: Workbook, plan: Plan) -> InputRefs:
    refs = InputRefs()
    s = plan.settings

    # settings
    ws = wb.create_sheet(SETTINGS_SHEET)
    header_row(ws, 1, ["Key", "Value", "Description"])
    values = s.model_dump()
    for i, (key, desc, fmt) in enumerate(SETTINGS_ROWS, start=2):
        ws.cell(i, 1, key)
        c = ws.cell(i, 2, values[key])
        if fmt:
            c.number_format = fmt
        ws.cell(i, 3, desc)
        refs.settings[key] = ref(SETTINGS_SHEET, f"B{i}")
    _table(ws, "tblSettings", 3, len(SETTINGS_ROWS))
    widths(ws, {"A": 20, "B": 22, "C": 62})

    # cost stacks
    flat = [(st, order, line) for st in plan.stacks for order, line in enumerate(st.lines, start=1)]
    cols: list[Col] = [
        ("Stack ID", lambda x: x[0].id, None), ("Stack name", lambda x: x[0].name, None),
        ("Order", lambda x: x[1], None), ("Line ID", lambda x: x[2].id, None),
        ("Label", lambda x: x[2].label, None), ("Type", lambda x: x[2].type.value, None),
        ("Rate", lambda x: x[2].rate, PCT), ("Base", lambda x: x[2].base, None),
        ("Amount", lambda x: x[2].amount, MONEY), ("Amount currency", lambda x: x[2].amount_currency, None),
        ("Note", lambda x: x[2].note, None),
        ("Cost of goods line", lambda x: "yes" if x[0].cogs_line == x[2].id else None, None)]
    rows = _write_rows(wb, STACK_SHEET, "tblCostStack", cols, flat,
                       {"A": 12, "B": 18, "C": 7, "D": 14, "E": 44, "F": 14, "G": 9, "H": 12, "I": 11,
                        "J": 10, "K": 34, "L": 12})
    for (st, _, line), r in zip(flat, rows):
        refs.lines[(st.id, line.id)] = {"rate": ref(STACK_SHEET, f"G{r}"), "amount": ref(STACK_SHEET, f"I{r}")}
    ws = wb[STACK_SHEET]
    types = DataValidation(type="list", formula1='"' + ",".join(t.value for t in LineType) + '"')
    ws.add_data_validation(types)
    types.add(f"F2:F{max(len(flat) + 1, 2)}")

    # products
    attr_names = sorted({k for sku in plan.skus for k in sku.attributes})
    cols = [
        ("SKU ID", lambda x: x.id, None), ("Name", lambda x: x.name, None), ("Size", lambda x: x.size, None),
        ("Size label", lambda x: x.size_label, None), ("Unit", lambda x: x.unit, None),
        ("Brand price", lambda x: x.fob, MONEY), ("Target retail (IDR)", lambda x: x.target_retail, IDR),
        ("Rule competitor", lambda x: x.pricing_rule.competitor if x.pricing_rule else None, None),
        ("Rule discount", lambda x: x.pricing_rule.discount if x.pricing_rule else None, PCT),
        ("Cost stack", lambda x: x.stack, None), ("Opening units", lambda x: x.opening_units, "#,##0"),
        ("MOQ", lambda x: x.moq, "#,##0"), ("Lead time (months)", lambda x: x.lead_time_months, "0"),
    ] + [(a, (lambda name: lambda x: x.attributes.get(name))(a), None) for a in attr_names]
    rows = _write_rows(wb, PRODUCTS_SHEET, "tblProducts", cols, plan.skus,
                       {"A": 14, "B": 42, "C": 8, "D": 10, "E": 8, "F": 11, "G": 15, "H": 15, "I": 11,
                        "J": 12, "K": 12, "L": 9, "M": 12})
    keys = {"name": "Name", "size": "Size", "size_label": "Size label", "unit": "Unit", "fob": "Brand price",
            "target": "Target retail (IDR)", "discount": "Rule discount", "opening": "Opening units",
            "moq": "MOQ", "lead": "Lead time (months)"}
    for sku, r in zip(plan.skus, rows):
        refs.skus[sku.id] = {k: ref(PRODUCTS_SHEET, f"{_col(cols, h)}{r}") for k, h in keys.items()}

    if plan.competitors:
        # competitors
        cols = [
            ("Competitor ID", lambda x: x.id, None), ("For product", lambda x: x.sku, None),
            ("Brand", lambda x: x.brand, None), ("Product", lambda x: x.product, None),
            ("Origin", lambda x: x.origin, None), ("Size", lambda x: x.size, None), ("Unit", lambda x: x.unit, None),
            ("Price (IDR)", lambda x: x.price_idr, IDR), ("Source", lambda x: x.source, None),
            ("Registration", lambda x: x.registration, None), ("Notes", lambda x: x.notes, None)]
        rows = _write_rows(wb, COMPETITORS_SHEET, "tblCompetitors", cols, plan.competitors,
                           {"A": 16, "B": 16, "C": 16, "D": 40, "E": 14, "F": 8, "G": 8, "H": 13, "I": 36,
                            "J": 16, "K": 30})
        for comp, r in zip(plan.competitors, rows):
            refs.competitors[comp.id] = {k: ref(COMPETITORS_SHEET, f"{_col(cols, h)}{r}") for k, h in
                                         {"brand": "Brand", "product": "Product", "origin": "Origin", "size": "Size",
                                          "unit": "Unit", "price": "Price (IDR)", "source": "Source",
                                          "registration": "Registration"}.items()}

    if plan.channels:
        # channels
        cols = [("Channel ID", lambda x: x.id, None), ("Name", lambda x: x.name, None), ("Kind", lambda x: x.kind, None),
                ("Margin", lambda x: x.margin, PCT), ("Package", lambda x: x.package, None)]
        rows = _write_rows(wb, CHANNELS_SHEET, "tblChannels", cols, plan.channels,
                           {"A": 14, "B": 28, "C": 12, "D": 10, "E": 14})
        for ch, r in zip(plan.channels, rows):
            refs.channels[ch.id] = {"name": ref(CHANNELS_SHEET, f"B{r}"), "margin": ref(CHANNELS_SHEET, f"D{r}")}

    if any(ch.margin_components for ch in plan.channels):
        fees = [(ch, f) for ch in plan.channels for f in ch.margin_components]
        cols = [("Channel ID", lambda x: x[0].id, None), ("Fee", lambda x: x[1].label, None),
                ("Rate", lambda x: x[1].rate, PCT), ("Plus VAT", lambda x: "yes" if x[1].plus_vat else "no", None)]
        rows = _write_rows(wb, MARGINS_SHEET, "tblChannelFees", cols, fees, {"A": 14, "B": 30, "C": 9, "D": 9})
        for (ch, _), r in zip(fees, rows):
            refs.fees.setdefault(ch.id, []).append({"rate": ref(MARGINS_SHEET, f"C{r}"),
                                                    "plus_vat": ref(MARGINS_SHEET, f"D{r}")})

    if plan.channels:
        # listings: a product × channel grid of weights; blank = not listed
        ws = wb.create_sheet(LISTINGS_SHEET)
        listing_channels = [ch for ch in plan.channels if not ch.package]
        header_row(ws, 1, ["SKU ID", "Name"] + [ch.id for ch in listing_channels])
        for r, sku in enumerate(plan.skus, start=2):
            ws.cell(r, 1, sku.id)
            ws.cell(r, 2, sku.name)
            for j, ch in enumerate(listing_channels, start=3):
                ws.cell(r, j, ch.listings.get(sku.id))
                refs.listings[(ch.id, sku.id)] = ref(LISTINGS_SHEET, f"{get_column_letter(j)}{r}")
        _table(ws, "tblListings", 2 + len(listing_channels), len(plan.skus))
        widths(ws, {"A": 14, "B": 42})
        ws.freeze_panes = "C2"

    if plan.drivers:
        # drivers: one row per factor, monthly values across
        labels = month_labels(s)
        ws = wb.create_sheet(DRIVERS_SHEET)
        header_row(ws, 1, DRIVER_COLS + labels)
        r = 2
        for d in plan.drivers:
            for f in d.factors:
                ws.append([d.id, d.channel, d.name, d.preset, f.name, f.kind, f.unit,
                           f.value if f.kind == "constant" else None]
                          + ([f.at(m) for m in range(s.months)] if f.kind == "monthly" else [None] * s.months))
                refs.factors.setdefault(d.id, []).append({
                    "name": f.name, "kind": ref(DRIVERS_SHEET, f"F{r}"), "constant": ref(DRIVERS_SHEET, f"H{r}"),
                    "months": [ref(DRIVERS_SHEET, f"{get_column_letter(9 + m)}{r}") for m in range(s.months)]})
                r += 1
        _table(ws, "tblDrivers", len(DRIVER_COLS) + s.months, r - 2)
        kinds = DataValidation(type="list", formula1='"constant,monthly"')
        ws.add_data_validation(kinds)
        kinds.add(f"F2:F{max(r - 1, 2)}")
        widths(ws, {"A": 14, "B": 14, "C": 24, "D": 18, "E": 30, "F": 10, "G": 12, "H": 10})
        ws.freeze_panes = "I2"

    if plan.drivers:
        # overrides
        cols = [("Channel ID", lambda x: x.channel, None), ("SKU ID", lambda x: x.sku, None),
                ("Month", lambda x: x.month, "0"), ("Units", lambda x: x.units, "#,##0")]
        _write_rows(wb, OVERRIDES_SHEET, "tblOverrides", cols, plan.overrides, {"A": 14, "B": 14, "C": 8, "D": 10})
        wb[OVERRIDES_SHEET].cell(1, 6, "Month 1 is the first month of the plan. Add rows freely; "
                                       "any (channel, product, month) listed here replaces the driver estimate.")
        last = OVERRIDE_ROWS + 1
        refs.overrides = {k: f"'{OVERRIDES_SHEET}'!${c}$2:${c}${last}"
                          for k, c in (("channel", "A"), ("sku", "B"), ("month", "C"), ("units", "D"))}

    if plan.packages:
        # packages (treatments): one row per component
        comps = [(p, c) for p in plan.packages for c in p.components]
        cols = [("Package ID", lambda x: x[0].id, None), ("Package name", lambda x: x[0].name, None),
                ("Price multiplier", lambda x: x[0].price_multiplier, "0.00"),
                ("Duration (min)", lambda x: x[0].duration_minutes, "0"), ("SKU ID", lambda x: x[1].sku, None),
                ("Usage per treatment", lambda x: x[1].usage, "0.00")]
        rows = _write_rows(wb, PACKAGES_SHEET, "tblPackages", cols, comps,
                           {"A": 14, "B": 34, "C": 15, "D": 13, "E": 14, "F": 18})
        for (p, c), r in zip(comps, rows):
            refs.package_usage[(p.id, c.sku)] = ref(PACKAGES_SHEET, f"F{r}")
            refs.packages.setdefault(p.id, {"multiplier": ref(PACKAGES_SHEET, f"C{r}")})

    if plan.containers:
        cols = [("Container", lambda x: x.name, None), ("Cost", lambda x: x.cost, MONEY), ("Units", lambda x: x.units, "#,##0")]
        rows = _write_rows(wb, CONTAINERS_SHEET, "tblContainers", cols, plan.containers, {"A": 26, "B": 12, "C": 12})
        refs.containers = [{"name": ref(CONTAINERS_SHEET, f"A{r}"), "cost": ref(CONTAINERS_SHEET, f"B{r}"),
                            "units": ref(CONTAINERS_SHEET, f"C{r}")} for r in rows]

    if plan.market:
        cols = [(h, (lambda k: lambda x: getattr(x, k))(h.lower()), None) for h in MARKET_COLS]
        _write_rows(wb, MARKET_SHEET, "tblMarket", cols, plan.market, {"A": 34, "B": 30, "C": 18, "D": 40})
    return refs


# ------------------------------------------------------------------ reading back

def _rows(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    out = []
    for r in rows[1:]:
        if all(_blank(v) is None for v in r):
            continue
        out.append({h: v for h, v in zip(headers, r) if h})
    return out


def _blank(v):
    return None if v is None or (isinstance(v, str) and not v.strip()) else v


def _str(v) -> str | None:
    v = _blank(v)
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return None if v is None else str(v)


def _num(v) -> float | None:
    v = _blank(v)
    return None if v is None else float(v)


def _yes(v) -> bool:
    return str(v).strip().lower() in ("yes", "y", "true", "1") if _blank(v) is not None else False


def read_inputs(path) -> Plan:
    """Rebuild a Plan from the `_Inputs` sheets of a workbook."""
    wb = load_workbook(path, data_only=True)
    required = (SETTINGS_SHEET, STACK_SHEET, PRODUCTS_SHEET)
    missing = [s for s in required if s not in wb.sheetnames]
    if missing:
        raise ValueError(f"workbook is missing input sheets: {', '.join(missing)}")

    def sheet(name):
        return _rows(wb[name]) if name in wb.sheetnames else []

    raw = {r["Key"]: _blank(r.get("Value")) for r in sheet(SETTINGS_SHEET) if r.get("Key")}
    settings = {k: v for k, v in raw.items() if v is not None}
    for key in ("rounding_step", "start_year", "start_month", "years"):
        if key in settings:
            settings[key] = int(settings[key])
    for key in ("hs_code", "origin", "incoterm", "brand_currency", "brand"):
        if key in settings:
            settings[key] = _str(settings[key])
    settings = Settings(**settings)

    stacks: dict[str, dict] = {}
    for r in sorted(sheet(STACK_SHEET), key=lambda r: (str(r["Stack ID"]), float(r["Order"] or 0))):
        st = stacks.setdefault(r["Stack ID"], {"id": r["Stack ID"], "name": r.get("Stack name") or r["Stack ID"],
                                               "lines": [], "cogs_line": None})
        st["lines"].append(CostLine(
            id=r["Line ID"], label=r["Label"], type=r["Type"], rate=_num(r.get("Rate")),
            base=_str(r.get("Base")), amount=_num(r.get("Amount")),
            amount_currency=_str(r.get("Amount currency")) or "brand", note=_str(r.get("Note")) or ""))
        if _yes(r.get("Cost of goods line")):
            st["cogs_line"] = r["Line ID"]

    skus = []
    for r in sheet(PRODUCTS_SHEET):
        rule = None
        if _blank(r.get("Rule competitor")) is not None:
            rule = PricingRule(competitor=_str(r["Rule competitor"]), discount=_num(r.get("Rule discount")) or 0.0)
        skus.append(SKU(
            id=_str(r["SKU ID"]), name=_str(r["Name"]), size=_num(r.get("Size")), size_label=_str(r.get("Size label")),
            unit=_str(r.get("Unit")), fob=_num(r.get("Brand price")), target_retail=_num(r.get("Target retail (IDR)")),
            pricing_rule=rule, stack=_str(r.get("Cost stack")) or "default",
            opening_units=_num(r.get("Opening units")) or 0, moq=_num(r.get("MOQ")),
            lead_time_months=int(_num(r.get("Lead time (months)")) or 0),
            attributes={k: _str(v) for k, v in r.items() if k not in PRODUCT_COLS and _blank(v) is not None}))

    competitors = [Competitor(
        id=_str(r["Competitor ID"]), sku=_str(r["For product"]), brand=_str(r.get("Brand")) or "",
        product=_str(r["Product"]), origin=_str(r.get("Origin")) or "", size=_num(r.get("Size")),
        unit=_str(r.get("Unit")), price_idr=_num(r["Price (IDR)"]), source=_str(r.get("Source")) or "",
        registration=_str(r.get("Registration")) or "", notes=_str(r.get("Notes")) or "")
        for r in sheet(COMPETITORS_SHEET)]

    fees: dict[str, list[MarginComponent]] = {}
    for r in sheet(MARGINS_SHEET):
        fees.setdefault(_str(r["Channel ID"]), []).append(
            MarginComponent(label=_str(r["Fee"]), rate=_num(r["Rate"]) or 0.0, plus_vat=_yes(r.get("Plus VAT"))))
    listings: dict[str, dict[str, float]] = {}
    if LISTINGS_SHEET in wb.sheetnames:
        for r in _rows(wb[LISTINGS_SHEET]):
            for k, v in r.items():
                if k not in ("SKU ID", "Name") and _blank(v) is not None:
                    listings.setdefault(k, {})[_str(r["SKU ID"])] = float(v)
    channels = [Channel(
        id=_str(r["Channel ID"]), name=_str(r["Name"]), kind=_str(r.get("Kind")) or "retail",
        margin=_num(r.get("Margin")), margin_components=fees.get(_str(r["Channel ID"]), []),
        listings=listings.get(_str(r["Channel ID"]), {}), package=_str(r.get("Package")))
        for r in sheet(CHANNELS_SHEET)]

    drivers: dict[str, dict] = {}
    if DRIVERS_SHEET in wb.sheetnames:
        ws = wb[DRIVERS_SHEET]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if _blank(row[0]) is None:
                continue
            d = drivers.setdefault(_str(row[0]), {"id": _str(row[0]), "channel": _str(row[1]), "name": _str(row[2]),
                                                  "preset": _str(row[3]) or "custom", "factors": []})
            kind = _str(row[5]) or "constant"
            months = [float(v) if _blank(v) is not None else 0.0 for v in row[8:8 + settings.months]]
            d["factors"].append(Factor(name=_str(row[4]), kind=kind, unit=_str(row[6]) or "",
                                       value=_num(row[7]) if kind == "constant" else None,
                                       values=_trim(months) if kind == "monthly" else []))

    overrides = [Override(channel=_str(r["Channel ID"]), sku=_str(r["SKU ID"]), month=int(r["Month"]),
                          units=float(r["Units"])) for r in sheet(OVERRIDES_SHEET)
                 if _blank(r.get("Units")) is not None]

    packages: dict[str, dict] = {}
    for r in sheet(PACKAGES_SHEET):
        p = packages.setdefault(_str(r["Package ID"]), {
            "id": _str(r["Package ID"]), "name": _str(r["Package name"]),
            "price_multiplier": _num(r.get("Price multiplier")) or 4.0,
            "duration_minutes": int(_num(r["Duration (min)"])) if _num(r.get("Duration (min)")) else None,
            "components": []})
        p["components"].append(PackageComponent(sku=_str(r["SKU ID"]), usage=_num(r["Usage per treatment"])))

    containers = [ContainerOption(name=_str(r["Container"]), cost=_num(r["Cost"]), units=_num(r["Units"]))
                  for r in sheet(CONTAINERS_SHEET)]
    market = [MarketRow(**{h.lower(): _str(r.get(h)) or "" for h in MARKET_COLS}) for r in sheet(MARKET_SHEET)]

    return Plan(settings=settings, stacks=[CostStack(**s) for s in stacks.values()], skus=skus,
                competitors=competitors, channels=channels, drivers=[Driver(**d) for d in drivers.values()],
                overrides=overrides, packages=[Package(**p) for p in packages.values()],
                containers=containers, market=market)


def _trim(values: list[float]) -> list[float]:
    """Drop trailing repeats so a monthly list round-trips to its shortest form."""
    out = list(values)
    while len(out) > 1 and out[-1] == out[-2]:
        out.pop()
    return out
