"""Price Structure, Competitor Analysis and Retail Prices sheets. Every number is a live formula."""

from __future__ import annotations

from dataclasses import dataclass

from openpyxl import Workbook

from .. import model
from ..schema import BASE_FOB, BASE_RUNNING, CostStack, LineType, Plan
from ..stack import linear_terms
from .inputs import InputRefs
from .style import (BOLD, IDR, INPUT, MONEY, PCT, SUBTITLE, SUBTOTAL_FILL, TITLE, TOP_RULE, header_row,
                    ref, widths)

PRICE_SHEET = "Price Structure"
COMPETITOR_SHEET = "Competitor Analysis"
RETAIL_SHEET = "Retail Prices"
EXAMPLE_FOB = 100


@dataclass
class Expected:
    """A cell whose recalculated value must match the Python engine."""
    sheet: str
    coord: str
    value: float
    label: str


@dataclass
class StackRefs:
    a: str               # retail coefficient: shelf price per 1 unit of brand price
    b: str | None        # fixed costs per unit in the shelf price, brand currency
    cogs_a: str | None   # cost-of-goods subtotal per 1 unit of brand price
    cogs_b: str | None


@dataclass
class SkuRefs:
    name: str
    forecast_price: str  # IDR incl. VAT used for sales
    cost_fob: str        # brand price used for purchases and cost of goods
    landed: str          # cost of goods per unit, IDR


def round_step(expr: str, step: str) -> str:
    return f"IF({step}=0,{expr},ROUNDUP(({expr})/{step},0)*{step})"


# ------------------------------------------------------------------ price structure

def write_price_structure(wb: Workbook, plan: Plan, refs: InputRefs,
                          expected: list[Expected]) -> dict[str, StackRefs]:
    ws = wb.create_sheet(PRICE_SHEET)
    cur = plan.settings.brand_currency
    widths(ws, {"A": 2, "B": 50, "C": 10, "D": 28, "E": 18, "F": 18})
    out: dict[str, StackRefs] = {}
    row = 1
    for stack in plan.stacks:
        has_fixed = any(line.type is LineType.FIXED for line in stack.lines)
        cols = ["E", "F"] if has_fixed else ["E"]

        ws.cell(row, 2, f"Price structure – {stack.name}").font = TITLE
        ws.cell(row + 1, 2, f"{plan.settings.incoterm} · amounts in {cur} · rates come from the "
                            f"'_Inputs Cost Stack' sheet").font = SUBTITLE
        row += 3
        heads = ["Line", "Rate", "Applied to", f"At brand price {EXAMPLE_FOB} ({cur})"]
        if has_fixed:
            heads.append(f"Fixed costs only ({cur})")
        header_row(ws, row, heads, start_col=2)
        row += 1

        fob_row = row
        ws.cell(row, 2, f"Brand price ({plan.settings.incoterm})").font = BOLD
        ws.cell(row, 5, EXAMPLE_FOB).number_format = MONEY
        if has_fixed:
            ws.cell(row, 6, 0).number_format = MONEY
        row += 1

        line_rows: dict[str, int] = {}
        contrib = [fob_row]  # rows summed into the running total since the last subtotal
        pending: list[int] = []
        labels = {line.id: line.label for line in stack.lines}
        next_sub = _next_subtotals(stack)
        for line in stack.lines:
            r = row
            line_rows[line.id] = r
            ws.cell(r, 2, line.label)
            lref = refs.lines[(stack.id, line.id)]
            if line.type in (LineType.PCT_OF_BASE, LineType.PCT_OF_RESULT):
                c = ws.cell(r, 3, f"={lref['rate']}")
                c.number_format, c.fill = PCT, INPUT

            for col in cols:
                formula = _formula(line, col, r, fob_row, contrib, pending, line_rows, lref, refs)
                if formula:
                    ws[f"{col}{r}"] = formula
                    ws[f"{col}{r}"].number_format = MONEY

            if line.type is LineType.PCT_OF_BASE:
                ws.cell(r, 4, {BASE_FOB: "brand price", BASE_RUNNING: "running total", None: "running total"}
                        .get(line.base, labels.get(line.base, line.base)))
                contrib.append(r)
            elif line.type is LineType.FIXED:
                ws.cell(r, 4, "fixed amount" + (" (IDR, converted)" if line.amount_currency == "idr" else ""))
                contrib.append(r)
            elif line.type is LineType.PCT_OF_RESULT:
                ws.cell(r, 4, f"share of {labels[next_sub[line.id]]}")
                pending.append(r)
            else:
                for p in pending:  # '% of resulting price' rows are a share of this subtotal
                    for col in cols:
                        ws[f"{col}{p}"] = f"=C{p}*{col}{r}"
                        ws[f"{col}{p}"].number_format = MONEY
                for col in range(2, 4 + len(cols) + 1):
                    ws.cell(r, col).fill = SUBTOTAL_FILL
                    ws.cell(r, col).border = TOP_RULE
                ws.cell(r, 2).font = BOLD
                for col in cols:
                    ws[f"{col}{r}"].font = BOLD
                contrib, pending = [r], []
            row += 1

        shelf = line_rows[stack.shelf.id]
        row += 1
        ws.cell(row, 2, f"Retail price incl. VAT, IDR (at brand price {EXAMPLE_FOB})")
        ws.cell(row, 5, f"=E{shelf}*{refs.fx}").number_format = IDR
        row += 1
        ws.cell(row, 2, "Retail coefficient (shelf price ÷ brand price)").font = BOLD
        c = ws.cell(row, 5, f"=(E{shelf}-F{shelf})/{EXAMPLE_FOB}" if has_fixed else f"=E{shelf}/{EXAMPLE_FOB}")
        c.number_format, c.font = "0.0000", BOLD
        a_ref = ref(PRICE_SHEET, f"E{row}")
        fx = plan.settings.fx_rate
        expected.append(Expected(PRICE_SHEET, f"E{row}", linear_terms(stack, fx)[0],
                                 f"coefficient of '{stack.name}'"))
        b_ref = None
        if has_fixed:
            row += 1
            ws.cell(row, 2, f"Fixed costs per unit in the shelf price ({cur})")
            ws.cell(row, 5, f"=F{shelf}").number_format = MONEY
            b_ref = ref(PRICE_SHEET, f"E{row}")

        cogs_a = cogs_b = None
        if stack.cogs_line:
            cl = line_rows[stack.cogs_line]
            row += 1
            ws.cell(row, 2, f"Cost of goods per 1 {cur} of brand price ({labels[stack.cogs_line]})")
            c = ws.cell(row, 5, f"=(E{cl}-F{cl})/{EXAMPLE_FOB}" if has_fixed else f"=E{cl}/{EXAMPLE_FOB}")
            c.number_format = "0.0000"
            cogs_a = ref(PRICE_SHEET, f"E{row}")
            if has_fixed:
                row += 1
                ws.cell(row, 2, f"Fixed costs per unit in cost of goods ({cur})")
                ws.cell(row, 5, f"=F{cl}").number_format = MONEY
                cogs_b = ref(PRICE_SHEET, f"E{row}")
        out[stack.id] = StackRefs(a=a_ref, b=b_ref, cogs_a=cogs_a, cogs_b=cogs_b)
        row += 3
    return out


def _formula(line, col, r, fob_row, contrib, pending, line_rows, lref, refs) -> str | None:
    def cells(rows):
        return "+".join(f"{col}{x}" for x in rows)

    if line.type is LineType.PCT_OF_BASE:
        if line.base == BASE_FOB:
            base = f"{col}{fob_row}"
        elif line.base in (None, BASE_RUNNING):
            base = f"({cells(contrib)})"
        else:
            base = f"{col}{line_rows[line.base]}"
        return f"=C{r}*{base}"
    if line.type is LineType.FIXED:
        if col == "E" or col == "F":
            amount = lref["amount"]
            return f"={amount}/{refs.fx}" if line.amount_currency == "idr" else f"={amount}"
    if line.type is LineType.PCT_OF_RESULT:
        return None  # written when the subtotal row is reached
    if line.type is LineType.SUBTOTAL:
        running = cells(contrib)
        if pending:
            return f"=({running})/(1-({'+'.join(f'C{p}' for p in pending)}))"
        return f"={running}"
    return None


def _next_subtotals(stack: CostStack) -> dict[str, str]:
    out, waiting = {}, []
    for line in stack.lines:
        if line.type is LineType.PCT_OF_RESULT:
            waiting.append(line.id)
        elif line.type is LineType.SUBTOTAL:
            out.update({w: line.id for w in waiting})
            waiting = []
    return out


# ------------------------------------------------------------------ competitors

def write_competitors(wb: Workbook, plan: Plan, refs: InputRefs, expected: list[Expected]) -> dict[str, str]:
    """Returns competitor id → size-matched price cell."""
    if not plan.competitors:
        return {}
    ws = wb.create_sheet(COMPETITOR_SHEET)
    s = plan.settings
    ws.cell(1, 2, f"Competitor analysis – {s.brand}").font = TITLE
    ws.cell(2, 2, "Price at our size = competitor price ÷ competitor size × our size, rounded up per the "
                  "rounding setting (only when both sizes use the same unit).").font = SUBTITLE
    heads = ["Brand", "Product", "Origin", "Size", "Unit", "Price (IDR)", "Price per unit (IDR)",
             "Price at our size (IDR)", f"Price at our size ({s.brand_currency})", "Source", "Registration"]
    out: dict[str, str] = {}
    row = 4
    for sku in plan.skus:
        comps = [c for c in plan.competitors if c.sku == sku.id]
        if not comps:
            continue
        sr = refs.skus[sku.id]
        ws.cell(row, 2, f"={sr['name']}").font = BOLD
        rule = sku.pricing_rule
        if rule:
            comp = plan.competitor(rule.competitor)
            direction = "below" if rule.discount >= 0 else "above"
            ws.cell(row, 3, f"Rule: {abs(rule.discount):.0%} {direction} {comp.brand} {comp.product}".strip())
        ws.cell(row, 5, f'=IF({sr["size"]}="","",{sr["size"]})')
        ws.cell(row, 6, f'=IF({sr["unit"]}="","",{sr["unit"]})')
        for col in range(2, 13):
            ws.cell(row, col).fill = SUBTOTAL_FILL
        row += 1
        header_row(ws, row, heads, start_col=2)
        row += 1
        for comp in comps:
            cr = refs.competitors[comp.id]
            for col, key in ((2, "brand"), (3, "product"), (4, "origin"), (10, "source"), (11, "registration")):
                ws.cell(row, col + (1 if col >= 10 else 0), f'=IF({cr[key]}="","",{cr[key]})')
            ws.cell(row, 5, f'=IF({cr["size"]}="","",{cr["size"]})')
            ws.cell(row, 6, f'=IF({cr["unit"]}="","",{cr["unit"]})')
            ws.cell(row, 7, f"={cr['price']}").number_format = IDR
            ws.cell(row, 8, f'=IF({cr["size"]}="","",{cr["price"]}/{cr["size"]})').number_format = IDR
            same = f'AND({cr["size"]}<>"",{sr["size"]}<>"",LOWER(TRIM({cr["unit"]}))=LOWER(TRIM({sr["unit"]})))'
            matched = round_step(f'{cr["price"]}/{cr["size"]}*{sr["size"]}', refs.step)
            c = ws.cell(row, 9, f"=IF({same},{matched},{cr['price']})")
            c.number_format, c.font = IDR, BOLD
            ws.cell(row, 10, f"=I{row}/{refs.fx}").number_format = MONEY
            out[comp.id] = ref(COMPETITOR_SHEET, f"I{row}")
            expected.append(Expected(COMPETITOR_SHEET, f"I{row}", model.size_matched_price(plan, comp),
                                     f"size-matched price of '{comp.product}'"))
            row += 1
        row += 1
    widths(ws, {"A": 2, "B": 18, "C": 40, "D": 14, "E": 8, "F": 8, "G": 13, "H": 13, "I": 15, "J": 14,
                "K": 36, "L": 16})
    return out


# ------------------------------------------------------------------ retail prices

def write_retail_prices(wb: Workbook, plan: Plan, refs: InputRefs, stacks: dict[str, StackRefs],
                        matched: dict[str, str], expected: list[Expected]) -> dict[str, SkuRefs]:
    ws = wb.create_sheet(RETAIL_SHEET)
    s = plan.settings
    cur = s.brand_currency
    ws.cell(1, 2, f"Retail prices – {s.brand}").font = TITLE
    ws.cell(2, 2, "Proposed retail = typed target, else the competitor rule. Forecast price = proposed retail, "
                  "else the retail price at the current brand price.").font = SUBTITLE
    heads = ["No", "Product", "Size", "Unit", f"Brand price ({cur})", "Cost stack", "Coefficient",
             "Retail at brand price (IDR)", "Proposed retail (IDR)", f"Brand price to ask for ({cur})",
             "vs current brand price", "Forecast price (IDR)", f"Cost basis ({cur})", "Cost of goods per unit (IDR)"]
    header_row(ws, 4, heads)
    stack_names = {st.id: st.name for st in plan.stacks}
    out: dict[str, SkuRefs] = {}
    for i, sku in enumerate(plan.skus, start=1):
        r = 4 + i
        sr, st = refs.skus[sku.id], stacks[sku.stack]
        ws.cell(r, 1, i)
        ws.cell(r, 2, f"={sr['name']}")
        ws.cell(r, 3, f'=IF({sr["size_label"]}<>"",{sr["size_label"]},IF({sr["size"]}="","",{sr["size"]}))')
        ws.cell(r, 4, f'=IF({sr["unit"]}="","",{sr["unit"]})')
        c = ws.cell(r, 5, f'=IF({sr["fob"]}="","",{sr["fob"]})')
        c.number_format, c.fill = MONEY, INPUT
        ws.cell(r, 6, stack_names[sku.stack])
        ws.cell(r, 7, f"={st.a}").number_format = "0.0000"
        fixed = f"+{st.b}" if st.b else ""
        ws.cell(r, 8, f'=IF(E{r}="","",{round_step(f"(E{r}*{st.a}{fixed})*{refs.fx}", refs.step)})'
                ).number_format = IDR

        target = f'{sr["target"]}'
        if sku.pricing_rule:
            rule = round_step(f'{matched[sku.pricing_rule.competitor]}*(1-{sr["discount"]})', refs.step)
            proposed = f'=IF({target}<>"",{target},{rule})'
        else:
            proposed = f'=IF({target}="","",{target})'
        c = ws.cell(r, 9, proposed)
        c.number_format, c.fill = IDR, INPUT
        b_part = f"-{st.b}" if st.b else ""
        ws.cell(r, 10, f'=IF(I{r}="","",(I{r}/{refs.fx}{b_part})/{st.a})').number_format = MONEY
        ws.cell(r, 11, f'=IF(OR(E{r}="",J{r}=""),"",J{r}/E{r}-1)').number_format = PCT
        c = ws.cell(r, 12, f'=IF(I{r}<>"",I{r},H{r})')
        c.number_format, c.font = IDR, BOLD
        ws.cell(r, 13, f'=IF(E{r}<>"",E{r},J{r})').number_format = MONEY
        if st.cogs_a:
            cb = f"+{st.cogs_b}" if st.cogs_b else ""
            ws.cell(r, 14, f'=IF(M{r}="",0,(M{r}*{st.cogs_a}{cb})*{refs.fx})').number_format = IDR
        else:
            ws.cell(r, 14, 0).number_format = IDR
        out[sku.id] = SkuRefs(name=ref(RETAIL_SHEET, f"B{r}"), forecast_price=ref(RETAIL_SHEET, f"L{r}"),
                              cost_fob=ref(RETAIL_SHEET, f"M{r}"), landed=ref(RETAIL_SHEET, f"N{r}"))

        for coord, value, label in (
                (f"H{r}", model.retail_from_fob(plan, sku), "retail price"),
                (f"I{r}", model.proposed_retail(plan, sku), "proposed retail"),
                (f"J{r}", model.implied_brand_price(plan, sku), "brand price to ask for"),
                (f"L{r}", model.forecast_price(plan, sku), "forecast price"),
                (f"N{r}", model.landed_cost_idr(plan, sku), "cost of goods")):
            if value is not None:
                expected.append(Expected(RETAIL_SHEET, coord, value, f"{label} of '{sku.name}'"))
    widths(ws, {"A": 5, "B": 42, "C": 9, "D": 7, "E": 12, "F": 14, "G": 11, "H": 15, "I": 15, "J": 15,
                "K": 12, "L": 15, "M": 12, "N": 15})
    ws.freeze_panes = "C5"
    return out
