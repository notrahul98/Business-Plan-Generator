"""Shared look for generated sheets."""

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import quote_sheetname

NAVY = "14213D"
TEAL = "0B6E6D"
SAND = "F1EEE6"
INPUT_FILL = "FFF4D6"  # pale yellow marks cells that come from _Inputs

TITLE = Font(name="Calibri", size=14, bold=True, color=NAVY)
SUBTITLE = Font(name="Calibri", size=10, italic=True, color="5F6673")
HEADER_FONT = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
BOLD = Font(name="Calibri", bold=True, color=NAVY)
SUBTOTAL_FILL = PatternFill("solid", fgColor=SAND)
INPUT = PatternFill("solid", fgColor=INPUT_FILL)
TOP_RULE = Border(top=Side(style="thin", color=NAVY))
WRAP = Alignment(wrap_text=True, vertical="top")

PCT = "0.0%"
MONEY = "#,##0.00"
IDR = "#,##0"


def ref(sheet: str, coord: str) -> str:
    """Absolute cross-sheet reference, e.g. '_Inputs Settings'!$B$6."""
    col = "".join(c for c in coord if c.isalpha())
    row = "".join(c for c in coord if c.isdigit())
    return f"{quote_sheetname(sheet)}!${col}${row}"


def header_row(ws, row: int, labels: list[str], start_col: int = 1) -> None:
    for i, label in enumerate(labels):
        c = ws.cell(row=row, column=start_col + i, value=label)
        c.font, c.fill = HEADER_FONT, HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")


def widths(ws, cols: dict[str, float]) -> None:
    for col, w in cols.items():
        ws.column_dimensions[col].width = w
