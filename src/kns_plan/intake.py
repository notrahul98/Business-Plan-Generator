"""Bring brand files into a plan: read any product/competitor/store list, guess the column mapping,
let the user correct it, then convert rows into plan objects. Unmapped columns become attributes."""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from .schema import SKU, Competitor, Plan, Settings, slugify

TEMPLATE_FIELDS = [  # header, explanation, example
    ("SKU", "Brand's product code (optional; used to recognise the product on later uploads)", "AB-001"),
    ("Product name", "Required", "Hydrating eye serum"),
    ("Size", "Number, or number with unit (15ml), or pack (31 x 4 g)", "15"),
    ("Unit", "ml, g, pcs, pads… (optional if the size includes it)", "ml"),
    ("Brand price", "Price from the brand per unit, in the brand's currency (FOB/EXW)", "9.50"),
    ("Target retail", "Optional: the shelf price you want, IDR incl. VAT", ""),
    ("MOQ", "Optional: minimum order quantity per product", ""),
]


@dataclass
class Table:
    headers: list[str]
    rows: list[dict[str, object]]
    sheet: str | None = None


@dataclass
class Converted:
    items: list = field(default_factory=list)
    skipped: list[str] = field(default_factory=list)  # human-readable reasons, one per skipped row


# ------------------------------------------------------------------ reading

def read_table(data: bytes, filename: str, sheet: str | None = None) -> Table:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        ws = wb[sheet] if sheet else wb.worksheets[0]
        return _from_grid([list(r) for r in ws.iter_rows(values_only=True)], ws.title)
    if name.endswith((".csv", ".txt", ".tsv")):
        text = _decode(data)
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t") if text.strip() else csv.excel
        return _from_grid([row for row in csv.reader(io.StringIO(text), dialect)])
    raise ValueError("upload an .xlsx or .csv file (save .xls files as .xlsx first)")


def sheet_names(data: bytes, filename: str) -> list[str]:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        return load_workbook(io.BytesIO(data), read_only=True).sheetnames
    return []


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _from_grid(grid: list[list], sheet: str | None = None) -> Table:
    """The header row is the first row with at least two filled cells."""
    grid = [r for r in grid if any(_text(v) for v in r)]
    start = next((i for i, r in enumerate(grid) if sum(1 for v in r if _text(v)) >= 2), 0)
    if not grid:
        return Table([], [], sheet)
    raw = grid[start]
    headers, seen = [], set()
    for i, h in enumerate(raw):
        h = _text(h) or f"Column {i + 1}"
        base, n = h, 2
        while h in seen:
            h, n = f"{base} ({n})", n + 1
        seen.add(h)
        headers.append(h)
    rows = []
    for r in grid[start + 1:]:
        row = {h: (r[i] if i < len(r) else None) for i, h in enumerate(headers)}
        if any(_text(v) for v in row.values()):
            rows.append(row)
    return Table(headers, rows, sheet)


def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


# ------------------------------------------------------------------ parsing values

def parse_number(v, idr: bool = False) -> float | None:
    """'$9.50' → 9.5, 'Rp 1.186.000' → 1186000, '9,50' → 9.5, '1,234.5' → 1234.5.
    With idr=True a single dot before exactly three digits is a thousands separator ('Rp 720.000')."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d,.\-]", "", str(v))
    if not s or s in "-.,":
        return None
    if "," in s and "." in s:
        dec = "," if s.rfind(",") > s.rfind(".") else "."
        s = s.replace("." if dec == "," else ",", "").replace(",", ".")
    elif "," in s:
        # groups of exactly three digits after each comma are thousands (1,186,000); otherwise decimal (9,50)
        parts = s.split(",")
        s = s.replace(",", "") if all(len(p) == 3 for p in parts[1:]) else s.replace(",", ".")
    elif s.count(".") > 1 or (idr and len(s.split(".")[-1]) == 3 and "." in s):
        s = s.replace(".", "")  # 1.186.000
    try:
        return float(s)
    except ValueError:
        return None


_SIZE = re.compile(r"^\s*(?P<a>\d+(?:[.,]\d+)?)\s*(?:[x×*]\s*(?P<b>\d+(?:[.,]\d+)?))?\s*(?P<unit>[a-zA-Z]+)?\s*$")


def parse_size(v) -> tuple[float | None, str | None, str | None]:
    """'15ml' → (15, 'ml', None); '31 x 4 g' → (124, 'g', '31 x 4'); 50 → (50, None, None)."""
    if v is None:
        return None, None, None
    if isinstance(v, (int, float)):
        return float(v), None, None
    m = _SIZE.match(str(v))
    if not m:
        return None, None, None
    a = float(m["a"].replace(",", "."))
    if m["b"]:
        b = float(m["b"].replace(",", "."))
        return a * b, m["unit"], f"{_text(a)} x {_text(b)}"
    return a, m["unit"], None


# ------------------------------------------------------------------ mapping

SYNONYMS: dict[str, dict[str, list[str]]] = {
    "products": {
        "id": ["sku id", "sku", "product code", "item code", "code", "article", "ref"],
        "name": ["name", "product name", "product", "sku name", "description", "item", "item name", "product description"],
        "size": ["size", "volume", "net content", "content", "net weight", "weight", "ml", "capacity"],
        "unit": ["unit", "uom", "unit of measure"],
        "fob": ["brand price", "fob", "exw", "fob price", "exw price", "unit price", "price", "cost", "supply price",
                "purchase price"],
        "target_retail": ["target retail", "rsp", "srp", "retail price", "recommended retail price", "proposed rsp"],
        "moq": ["moq", "minimum order", "min order"],
    },
    "competitors": {
        "sku": ["for product", "our product", "benchmark for", "our sku", "sku"],
        "brand": ["brand", "competitor brand", "competitor"],
        "product": ["product", "competitor product", "name", "product name", "description"],
        "origin": ["origin", "coo", "country of origin", "country"],
        "size": ["size", "volume", "content"],
        "unit": ["unit", "uom"],
        "price_idr": ["price (idr)", "price", "retail price", "rsp", "retail prices (idr)", "selling price"],
        "source": ["source", "source link", "link", "url", "retailer"],
        "registration": ["registration", "bpom", "bpom notification no.", "notification", "reg no"],
    },
    "stores": {
        "channel": ["channel", "retailer", "account", "chain", "distribution channel"],
        "store": ["store", "store name", "outlet", "location", "branch", "name"],
        "opening": ["opening", "opening month", "open", "start", "start month", "launch", "listing month", "month"],
    },
}
REQUIRED = {"products": ["name"], "competitors": ["product", "price_idr"], "stores": ["channel", "opening"]}


def _norm(h: str) -> str:
    return re.sub(r"[^a-z0-9()]+", " ", h.lower()).strip()


def guess_mapping(kind: str, headers: list[str]) -> dict[str, str | None]:
    """Field → header. Exact synonym matches first, then headers that start with a synonym."""
    mapping: dict[str, str | None] = {}
    used: set[str] = set()
    for exact in (True, False):
        for fld, words in SYNONYMS[kind].items():
            if mapping.get(fld):
                continue
            for w in words:
                hit = next((h for h in headers if h not in used and
                            (_norm(h) == w if exact else _norm(h).startswith(w + " ") or _norm(h).startswith(w))),
                           None)
                if hit:
                    mapping[fld] = hit
                    used.add(hit)
                    break
    return {fld: mapping.get(fld) for fld in SYNONYMS[kind]}


def to_skus(table: Table, mapping: dict[str, str | None], existing: set[str] | None = None) -> Converted:
    out = Converted()
    ids = set(existing or ())
    mapped = {h for h in mapping.values() if h}
    for n, row in enumerate(table.rows, start=1):
        get = lambda f: row.get(mapping[f]) if mapping.get(f) else None  # noqa: E731
        name = _text(get("name"))
        if not any(_text(get(f)) for f in mapping if f != "name") and not name:
            continue  # a numbered but otherwise empty line
        if not name:
            out.skipped.append(f"data row {n}: no product name")
            continue
        size, unit, label = parse_size(get("size"))
        unit = _text(get("unit")) or unit
        sku_id = slugify(_text(get("id")) or name, ids)
        ids.add(sku_id)
        attrs = {h: _text(v) for h, v in row.items() if h not in mapped and _text(v)}
        if mapping.get("id") and _text(get("id")):
            attrs.setdefault("Brand code", _text(get("id")))
        try:
            out.items.append(SKU(id=sku_id, name=name, size=size if size and size > 0 else None, size_label=label,
                                 unit=unit or None, fob=parse_number(get("fob")),
                                 target_retail=parse_number(get("target_retail"), idr=True) or None,
                                 moq=parse_number(get("moq")) or None, attributes=attrs))
        except ValueError as e:
            out.skipped.append(f"data row {n} ({name}): {e.errors()[0]['msg'] if hasattr(e, 'errors') else e}")
    return out


def to_competitors(table: Table, mapping: dict[str, str | None], plan: Plan) -> Converted:
    """`sku` may hold our product's id or name; rows that match no product are skipped."""
    out = Converted()
    by_name = {s.name.strip().lower(): s.id for s in plan.skus}
    sku_ids = {s.id for s in plan.skus}
    ids = {c.id for c in plan.competitors}
    for n, row in enumerate(table.rows, start=1):
        get = lambda f: row.get(mapping[f]) if mapping.get(f) else None  # noqa: E731
        product = _text(get("product"))
        price = parse_number(get("price_idr"), idr=True)
        ours = _text(get("sku"))
        sku_id = ours if ours in sku_ids else by_name.get(ours.lower())
        if not product or not price:
            out.skipped.append(f"data row {n}: needs a product and a price")
            continue
        if not sku_id:
            out.skipped.append(f"data row {n} ({product}): '{ours}' does not match any of our products")
            continue
        size, unit, _ = parse_size(get("size"))
        cid = slugify(f"{_text(get('brand'))} {product}", ids)
        ids.add(cid)
        out.items.append(Competitor(id=cid, sku=sku_id, brand=_text(get("brand")), product=product,
                                    origin=_text(get("origin")), size=size or None,
                                    unit=_text(get("unit")) or unit, price_idr=price, source=_text(get("source")),
                                    registration=_text(get("registration"))))
    return out


def stores_by_month(table: Table, mapping: dict[str, str | None], settings: Settings) -> dict[str, list[float]]:
    """Channel name → number of open stores in each plan month (stores stay open once opened)."""
    n = settings.months
    start = settings.start_year * 12 + settings.start_month - 1
    counts: dict[str, list[float]] = {}
    for row in table.rows:
        channel = _text(row.get(mapping["channel"])) if mapping.get("channel") else ""
        opened = _month_index(row.get(mapping["opening"]) if mapping.get("opening") else None, start)
        if not channel or opened is None:
            continue
        series = counts.setdefault(channel, [0.0] * n)
        for m in range(max(opened, 0), n):
            series[m] += 1
    return counts


def _month_index(v, start: int) -> int | None:
    """Plan month (0-based) for a date, '2027-03', 'Mar 2027' or a plan month number (1 = first)."""
    if isinstance(v, (datetime, date)):
        return v.year * 12 + v.month - 1 - start
    if isinstance(v, (int, float)):
        return int(v) - 1
    s = _text(v)
    if not s:
        return None
    if s.isdigit():
        return int(s) - 1
    for fmt in ("%Y-%m", "%Y/%m", "%b %Y", "%B %Y", "%m/%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            d = datetime.strptime(s, fmt)
            return d.year * 12 + d.month - 1 - start
        except ValueError:
            continue
    return None


# ------------------------------------------------------------------ KNS template

def write_template(path: Path | str | None = None) -> bytes:
    """The KNS product template: the columns the tool understands, with one example row."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    ws.append([f[0] for f in TEMPLATE_FIELDS])
    ws.append([f[2] for f in TEMPLATE_FIELDS])
    for i, f in enumerate(TEMPLATE_FIELDS, start=1):
        ws.cell(1, i).font = Font(bold=True)
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(14, len(f[0]) + 4)
    notes = wb.create_sheet("How to fill")
    notes.append(["Column", "What to enter"])
    for f in TEMPLATE_FIELDS:
        notes.append([f[0], f[1]])
    notes.append([])
    notes.append(["Any other column", "Kept as a product attribute and shown on listing sheets; not used in formulas."])
    notes.column_dimensions["A"].width, notes.column_dimensions["B"].width = 24, 90
    buf = io.BytesIO()
    wb.save(buf)
    if path:
        Path(path).write_bytes(buf.getvalue())
    return buf.getvalue()
